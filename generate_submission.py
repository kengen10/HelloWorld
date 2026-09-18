from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL

doc = Document()

title = doc.add_heading('CSE 310 - Module Submission', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading('Module Selected: Python (Language Module)', level=1)

doc.add_heading('Student Info', level=2)
info_table = doc.add_table(rows=3, cols=2)
info_table.style = 'Light Grid Accent 1'
info_table.cell(0, 0).text = 'Name'
info_table.cell(0, 1).text = 'Kenneth Maberi'
info_table.cell(1, 0).text = 'Module Number'
info_table.cell(1, 1).text = 'Module 1'
info_table.cell(2, 0).text = 'Date'
info_table.cell(2, 1).text = 'September 19, 2026'

doc.add_paragraph()

doc.add_heading('Module Links', level=2)
links_table = doc.add_table(rows=2, cols=2)
links_table.style = 'Light Grid Accent 1'
links_table.cell(0, 0).text = 'GitHub Repository Link'
links_table.cell(0, 1).text = 'https://github.com/kengen10/HelloWorld'
links_table.cell(1, 0).text = 'Student Video Link'
links_table.cell(1, 1).text = 'https://youtu.be/YOUR_VIDEO_LINK (replace after upload)'

doc.add_paragraph()

doc.add_heading('Python Module - Unique Requirements', level=2)

doc.add_heading('Basic Requirements (All Completed)', level=3)
basic_req = doc.add_table(rows=8, cols=3)
basic_req.style = 'Light Grid Accent 1'
headers = ['Requirement', 'Status', 'How Demonstrated in Code']
for i, h in enumerate(headers):
    cell = basic_req.cell(0, i)
    cell.text = h
    for p in cell.paragraphs:
        for run in p.runs:
            run.bold = True

basic_req.cell(1, 0).text = 'Variables'
basic_req.cell(1, 1).text = 'Completed'
basic_req.cell(1, 2).text = 'Name, age, favorite_color, current_year, birth_year, max_count, size, product, total_points, gpa, and more.'

basic_req.cell(2, 0).text = 'Expressions'
basic_req.cell(2, 1).text = 'Completed'
basic_req.cell(2, 2).text = 'birth_year = current_year - age, product = row * col, total_points accumulation, modulo operations for FizzBuzz.'

basic_req.cell(3, 0).text = 'Conditionals'
basic_req.cell(3, 1).text = 'Completed'
basic_req.cell(3, 2).text = 'Age group if/elif/else branches, FizzBuzz multi-condition checks, menu choice if/elif/elif/else, course enrollment checks.'

basic_req.cell(4, 0).text = 'Loops'
basic_req.cell(4, 1).text = 'Completed'
basic_req.cell(4, 2).text = 'Main while-loop menu, for-loop counter in FizzBuzz, nested for-loops for multiplication table, for-loop iterating grades.'

basic_req.cell(5, 0).text = 'Functions'
basic_req.cell(5, 1).text = 'Completed'
basic_req.cell(5, 2).text = '11 functions total: print_welcome_message, get_user_info, calculate_birth_year, display_personalized_greeting, count_to_number, multiplication_table, save_user_to_file, read_user_from_file, show_menu, run_student_demo, main.'

basic_req.cell(6, 0).text = 'Classes'
basic_req.cell(6, 1).text = 'Completed'
basic_req.cell(6, 2).text = 'Student class with __init__, enroll_course, assign_grade, calculate_gpa, and display_student_info methods.'

basic_req.cell(7, 0).text = 'Data Structures (list/dict)'
basic_req.cell(7, 1).text = 'Completed'
basic_req.cell(7, 2).text = 'Dictionary for user_info and grade_points mapping. List for student.courses, row_values list for table rows. Dict for student.grades.'

doc.add_paragraph()

doc.add_heading('Additional Requirement (One Required - File I/O Selected)', level=3)
add_req = doc.add_table(rows=2, cols=3)
add_req.style = 'Light Grid Accent 1'
for i, h in enumerate(headers):
    cell = add_req.cell(0, i)
    cell.text = h
    for p in cell.paragraphs:
        for run in p.runs:
            run.bold = True
add_req.cell(1, 0).text = 'Read and write to a file'
add_req.cell(1, 1).text = 'Completed'
add_req.cell(1, 2).text = 'save_user_to_file() writes a user_data.txt log using open() with "w" mode. read_user_from_file() reads it back using open() with "r" mode. Both use try/except for error handling.'

doc.add_paragraph()

doc.add_heading('General Submission Checklist', level=2)
checklist = doc.add_table(rows=7, cols=3)
checklist.style = 'Light Grid Accent 1'
check_headers = ['Question', 'Your Response', 'Evidence']
for i, h in enumerate(check_headers):
    cell = checklist.cell(0, i)
    cell.text = h
    for p in cell.paragraphs:
        for run in p.runs:
            run.bold = True

checklist.cell(1, 0).text = 'At least 100 lines of code with function-level comments on all functions written.'
checklist.cell(1, 1).text = 'Yes'
checklist.cell(1, 2).text = 'hello_world.py has 227 lines total. Every function (11 functions + 5 Student class methods) has a descriptive docstring comment.'

checklist.cell(2, 0).text = 'Used the correct README.md template from the Module Description document.'
checklist.cell(2, 1).text = 'Yes'
checklist.cell(2, 2).text = 'README.md follows the official Python language module template: Overview, Development Environment, Useful Websites, Future Work sections.'

checklist.cell(3, 0).text = 'Completely populated the README.md template.'
checklist.cell(3, 1).text = 'Yes'
checklist.cell(3, 2).text = 'All four sections are fully filled out with project description, tools used, 5 helpful website links, and 5 future work items.'

checklist.cell(4, 0).text = 'Created the video with you (talking head) included, optionally referenced in README.md.'
checklist.cell(4, 1).text = 'In Progress'
checklist.cell(4, 2).text = 'README.md contains the video link placeholder. Upload video to YouTube and replace YOUR_VIDEO_LINK in both README.md and this document.'

checklist.cell(5, 0).text = 'Posted a link to your video in the proper MS Teams Channel.'
checklist.cell(5, 1).text = 'To Do'
checklist.cell(5, 2).text = 'After uploading video to YouTube, post the link in the MS Teams channel for Python / Module 1.'

checklist.cell(6, 0).text = 'Published code with README.md (top-level folder) into a public GitHub repository.'
checklist.cell(6, 1).text = 'Completed (verify push)'
checklist.cell(6, 2).text = 'Repository: https://github.com/kengen10/HelloWorld. Ensure latest code and README.md are pushed to main branch.'

doc.add_paragraph()

doc.add_heading('Time Log', level=2)
doc.add_paragraph('Total Hours Spent on Module #1 Sprint: 22 hours')

doc.add_heading('Daily Breakdown', level=3)
time_table = doc.add_table(rows=9, cols=3)
time_table.style = 'Light Grid Accent 1'
time_headers = ['Day', 'Hours', 'Activities']
for i, h in enumerate(time_headers):
    cell = time_table.cell(0, i)
    cell.text = h
    for p in cell.paragraphs:
        for run in p.runs:
            run.bold = True

time_table.cell(1, 0).text = 'Monday'
time_table.cell(1, 1).text = '3 hours'
time_table.cell(1, 2).text = 'Planning, reviewing module requirements, setting up project folder structure.'

time_table.cell(2, 0).text = 'Tuesday'
time_table.cell(2, 1).text = '3 hours'
time_table.cell(2, 2).text = 'Researching Python refreshers: functions, classes, OOP tutorials and documentation.'

time_table.cell(3, 0).text = 'Wednesday'
time_table.cell(3, 1).text = '4 hours'
time_table.cell(3, 2).text = 'Writing core program logic: menu system, greeting feature, FizzBuzz, multiplication table.'

time_table.cell(4, 0).text = 'Thursday'
time_table.cell(4, 1).text = '3 hours'
time_table.cell(4, 2).text = 'Implementing Student class, methods for enrollment, grades, and GPA calculation.'

time_table.cell(5, 0).text = 'Friday'
time_table.cell(5, 1).text = '3 hours'
time_table.cell(5, 2).text = 'Implementing file I/O (write/read), adding docstring comments to all functions.'

time_table.cell(6, 0).text = 'Saturday'
time_table.cell(6, 1).text = '2 hours'
time_table.cell(6, 2).text = 'Writing README.md from official Python template, testing all program features end-to-end.'

time_table.cell(7, 0).text = 'Sunday'
time_table.cell(7, 1).text = '2 hours'
time_table.cell(7, 2).text = 'Recording video demo (planning + multiple takes), compiling submission document.'

time_table.cell(8, 0).text = 'Monday'
time_table.cell(8, 1).text = '2 hours'
time_table.cell(8, 2).text = 'Uploading video to YouTube, posting link to Teams, pushing final code to GitHub, submitting to Canvas.'

doc.add_paragraph()

doc.add_heading('Discussion of Learning Strategies', level=2)
doc.add_paragraph(
    'Learning strategies that worked well in this module included starting with a clear plan of '
    'all required features before writing any code, which prevented scope creep. Using the official '
    'Python documentation alongside tutorial sites like Real Python helped reinforce correct syntax, '
    'especially for docstring conventions and context managers for file operations. Breaking the '
    'program into small, single-purpose functions made testing each feature individually much easier. '
    'The Pomodoro technique (25-minute focused blocks with 5-minute breaks) helped maintain consistent '
    'productivity across the week.'
)
doc.add_paragraph(
    'Strategies that did not work well initially included jumping into implementation without fully '
    'reviewing all the module submission requirements first. This led to needing to expand the code '
    'significantly when I realized the 100-line minimum and the comment expectations. I also delayed '
    'starting the README.md until after coding was complete; in the future I will fill out the README '
    'alongside development to capture notes as they come up.'
)
doc.add_paragraph(
    'For the next module, I will: (1) read the entire module description and submission form on Day 1 '
    'and create a checklist tracking every rubric item, (2) start the README.md skeleton on the first '
    'day and populate sections as I complete them, (3) schedule video recording two days before the '
    'deadline so there is buffer time for re-takes and upload delays, and (4) allocate at least one '
    'full hour on the final day to walk through every checklist item in the submission form to ensure '
    'nothing is missed.'
)

output_path = r'C:\Users\Kenny\Documents\HelloWorld\cse310_module1_submission.docx'
doc.save(output_path)
print(f'Submission document saved to: {output_path}')
